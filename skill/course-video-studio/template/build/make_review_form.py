# -*- coding: utf-8 -*-
"""Build the viewing-notes workbook: one row per note, lessons as a dropdown."""
import json, sys
sys.path.insert(0, ".")
from course import BUILD, LESSONS, ROOT, lesson_meta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

AUDIO_QC = BUILD / "qc" / "audio_qc.json"   # durations, when audio_qc.py has run
aud = {r["id"]: r for r in json.loads(AUDIO_QC.read_text(encoding="utf-8"))} if AUDIO_QC.exists() else {}
wb = Workbook()
ws = wb.active; ws.title = "הערות"; ws.sheet_view.rightToLeft = True
heads = ["שיעור", "זמן בסרטון (דקה:שנייה)", "מה ראיתי או שמעתי", "מה לתקן", "חומרה", "טופל"]
widths = [34, 20, 48, 48, 12, 8]
thin = Side(style="thin", color="000000"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c, (h, w) in enumerate(zip(heads, widths), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name="David", size=12, bold=True); cell.fill = PatternFill("solid", fgColor="D9D9D9")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2); cell.border = border
    ws.column_dimensions[get_column_letter(c)].width = w
names = [f'{lesson_meta(l[0])["num"]} {l[1]}' for l in LESSONS]
ref = wb.create_sheet("שיעורים"); ref.sheet_view.rightToLeft = True
for c, h in enumerate(["שיעור", "משך", "פרק"], 1):
    cell = ref.cell(row=1, column=c, value=h); cell.font = Font(name="David", size=12, bold=True); cell.fill = PatternFill("solid", fgColor="D9D9D9"); cell.border = border
for i, (lid, title, promise, m) in enumerate(LESSONS, 2):
    d = aud.get(lid, {}).get("dur", 0)
    for c, v in enumerate([names[i - 2], f"{int(d // 60)}:{int(d % 60):02d}", m["name"]], 1):
        cell = ref.cell(row=i, column=c, value=v); cell.font = Font(name="David", size=12); cell.border = border
        cell.alignment = Alignment(readingOrder=2, vertical="center")
ref.column_dimensions["A"].width = 44; ref.column_dimensions["B"].width = 8; ref.column_dimensions["C"].width = 26
dv = DataValidation(type="list", formula1=f"=שיעורים!$A$2:$A${len(LESSONS) + 1}", allow_blank=True)
sev = DataValidation(type="list", formula1='"חובה,מומלץ,רעיון"', allow_blank=True)
done = DataValidation(type="list", formula1='"כן,לא"', allow_blank=True)
ws.add_data_validation(dv); ws.add_data_validation(sev); ws.add_data_validation(done)
N = 200
for r in range(2, N + 2):
    for c in range(1, 7):
        cell = ws.cell(row=r, column=c); cell.font = Font(name="David", size=12); cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="top", readingOrder=2)
    ws.row_dimensions[r].height = 30
dv.add(f"A2:A{N + 1}"); sev.add(f"E2:E{N + 1}"); done.add(f"F2:F{N + 1}")
ws.freeze_panes = "A2"
out = ROOT / "05 - טופס הערות לצפייה.xlsx"
wb.save(out); print(out.name, out.stat().st_size)
